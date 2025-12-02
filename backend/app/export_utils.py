# /geoconverter/backend/app/export_utils.py
from typing import List
from io import BytesIO

from openpyxl import Workbook

from .schemas import Point


def _safe(v):
    return v if v is not None else ""


def build_catalog_xlsx(rows: List[Point], crs_from: str, crs_to: str) -> bytes:
    """Простой каталог координат по результатам пересчёта.

    Структура:
    - №
    - Наименование
    - Широта
    - Долгота
    - X
    - Y
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Каталог"

    # Шапка
    ws.append([f"Каталог координат ({crs_from} → {crs_to})"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)

    # Заголовки колонок
    ws.append(["№", "Наименование", "Широта", "Долгота", "X", "Y"])

    for idx, p in enumerate(rows, start=1):
        ws.append(
            [
                idx,
                _safe(getattr(p, "name", None)),
                _safe(p.lat),
                _safe(p.lon),
                _safe(p.x),
                _safe(p.y),
            ]
        )

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def build_statement_xlsx(
    rows_in: List[Point],
    rows_out: List[Point],
    crs_from: str,
    crs_to: str,
) -> bytes:
    """Ведомость пересчёта: исходные и полученные координаты в одной таблице."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ведомость"

    max_len = max(len(rows_in), len(rows_out)) if (rows_in or rows_out) else 0

    ws.append([f"Ведомость расчёта координат ({crs_from} → {crs_to})"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)

    ws.append(
        [
            "№",
            "Наименование",
            "Исх. широта",
            "Исх. долгота",
            "Исх. X",
            "Исх. Y",
            "Рез. широта",
            "Рез. долгота",
            "Рез. X",
            "Рез. Y",
        ]
    )

    for idx in range(max_len):
        p_in = rows_in[idx] if idx < len(rows_in) else Point()
        p_out = rows_out[idx] if idx < len(rows_out) else Point()

        ws.append(
            [
                idx + 1,
                _safe(getattr(p_out, "name", None) or getattr(p_in, "name", None)),
                _safe(p_in.lat),
                _safe(p_in.lon),
                _safe(p_in.x),
                _safe(p_in.y),
                _safe(p_out.lat),
                _safe(p_out.lon),
                _safe(p_out.x),
                _safe(p_out.y),
            ]
        )

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
