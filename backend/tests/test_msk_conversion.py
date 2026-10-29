import importlib
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook

from app.msk_registry import DEFINITIONS, MSK_ITEMS, READY_MSK_CRS
from app.schemas import ConvertRequest, Point


MSK47 = 'MSK:GEOPROJ:47:2'
EXPECTED_MSK = {
    f'MSK:GEOPROJ:{region}:{zone}'
    for region, zones in [('02', 2), ('16', 3), ('47', 3), ('50', 2),
                          ('52', 3), ('63', 2), ('1964', 1)]
    for zone in range(1, zones + 1)
}


class MskTests(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        cls.uploads = tempfile.TemporaryDirectory()
        with patch.dict(os.environ, {'UPLOAD_DIR': cls.uploads.name}):
            cls.api = importlib.import_module('app.main')

    @classmethod
    def tearDownClass(cls):
        cls.uploads.cleanup()

    def setUp(self):
        upload_patch = patch.object(self.api, 'UPLOAD_DIR', Path(self.uploads.name))
        upload_patch.start()
        self.addCleanup(upload_patch.stop)

    def convert(self, source, target, points):
        return self.api.api_convert(ConvertRequest(crs_from=source, crs_to=target, points=points))

    async def test_registry_matches_executable_definitions(self):
        self.assertEqual(len(MSK_ITEMS), 259)
        self.assertEqual(len({item.code for item in MSK_ITEMS}), 259)
        self.assertEqual(set(READY_MSK_CRS), EXPECTED_MSK)
        for item in MSK_ITEMS:
            self.assertEqual(item.status == 'ready', item.code in READY_MSK_CRS)
            self.assertTrue(item.region)
            self.assertTrue(item.description)

    async def test_geoproj_reference_points_and_axis_convention(self):
        for definition in DEFINITIONS:
            for point in definition['control_points'] + definition.get('additional_control_points', []):
                with self.subTest(code=definition['code'], lat=point['lat']):
                    result = self.convert('EPSG:4326', definition['code'],
                                          [Point(lat=point['lat'], lon=point['lon'])])
                    actual = result.points[0]
                    tolerance = definition['verification_tolerance_metres']
                    self.assertAlmostEqual(actual.x, point['easting'], delta=tolerance)
                    self.assertAlmostEqual(actual.y, point['northing'], delta=tolerance)
                    self.assertFalse(result.meta['allow_ballpark'])
                    self.assertTrue(result.meta['operations'])
                    self.assertTrue(result.meta['regional_definitions'])
                    restored = self.convert(definition['code'], 'EPSG:4326', result.points).points[0]
                    self.assertAlmostEqual(restored.lat, point['lat'], places=6)
                    self.assertAlmostEqual(restored.lon, point['lon'], places=6)

    async def test_every_projection_origin_and_published_datum(self):
        from pyproj import CRS, Transformer
        for definition in DEFINITIONS:
            code = definition['code']
            with self.subTest(code=code):
                result = self.convert('SK42_GEOG', code, [Point(
                    lat=definition['latitude_natural_origin'],
                    lon=definition['longitude_natural_origin'])]).points[0]
                self.assertAlmostEqual(result.x, definition['false_easting'], places=6)
                self.assertAlmostEqual(result.y, definition['false_northing'], places=6)
                # Source explicitly supplies a datum shift. Verify that naming
                # Pulkovo 1942 has preserved the published WGS84 behaviour.
                published = CRS.from_proj4(definition['published_proj4'])
                transformer = Transformer.from_crs(4326, published, always_xy=True, allow_ballpark=False)
                for point in definition['control_points']:
                    x, y = transformer.transform(point['lon'], point['lat'], errcheck=True)
                    expected = self.convert('EPSG:4326', code,
                        [Point(lat=point['lat'], lon=point['lon'])]).points[0]
                    self.assertAlmostEqual(expected.x, x, delta=.01)
                    self.assertAlmostEqual(expected.y, y, delta=.01)

    async def test_msk50_cross_zone_and_new_sk63_paths(self):
        start = self.convert('EPSG:4326', 'MSK:GEOPROJ:50:1', [Point(lat=55.75, lon=37.6)])
        for target in ('MSK:GEOPROJ:50:2', 'MSK:GEOPROJ:16:1', 'SK63:GEOPROJ:P:2',
                       'SK63:GEOPROJ:C:1', 'SK63:GEOPROJ:X:4'):
            with self.subTest(target=target):
                result = self.convert('MSK:GEOPROJ:50:1', target, start.points)
                back = self.convert(target, 'MSK:GEOPROJ:50:1', result.points).points[0]
                self.assertAlmostEqual(back.x, start.points[0].x, delta=.02)
                self.assertAlmostEqual(back.y, start.points[0].y, delta=.02)

    async def test_projection_origin_and_cross_systems(self):
        origin = self.convert('SK42_GEOG', MSK47, [Point(lat=.1, lon=30.95)]).points[0]
        self.assertAlmostEqual(origin.x, 2250000, places=6)
        self.assertAlmostEqual(origin.y, -6200000, places=6)
        original = self.convert('EPSG:4326', MSK47, [Point(lat=59.95, lon=30.3)])
        for target in ('SK63:EPSG:3351', 'GK_SK42', 'GK_SK95', 'GK_GSK2011', 'SK95_GEOG'):
            with self.subTest(target=target):
                converted = self.convert(MSK47, target, original.points)
                restored = self.convert(target, MSK47, converted.points).points[0]
                self.assertAlmostEqual(restored.x, original.points[0].x, delta=.02)
                self.assertAlmostEqual(restored.y, original.points[0].y, delta=.02)

    async def test_pending_unknown_and_invalid_points_are_rejected(self):
        for code in [item.code for item in MSK_ITEMS if item.status == 'dev'] + ['MSK:UNKNOWN', 'MSK-50']:
            for source, target in ((code, 'EPSG:4326'), ('EPSG:4326', code), (code, 'SK63:EPSG:3351')):
                with self.subTest(source=source, target=target), self.assertRaises(HTTPException) as error:
                    self.convert(source, target, [])
                self.assertEqual(error.exception.status_code, 400)
        for points in ([Point(lat=59.95, lon=30.3)], [Point(x=1e100, y=1e100)]):
            with self.assertRaises(HTTPException) as error:
                self.convert(MSK47, 'EPSG:4326', points)
            self.assertEqual(error.exception.status_code, 400)

    async def test_100_points_import_convert_and_export(self):
        code = 'MSK:GEOPROJ:50:2'
        original = [Point(lat=55.75 + i / 10000, lon=37.6 + i / 10000) for i in range(100)]
        projected = self.convert('EPSG:4326', code, original).points
        workbook = Workbook()
        workbook.active.append(['X (м)', 'Y (м)'])
        for point in projected:
            workbook.active.append([point.x, point.y])
        content = io.BytesIO()
        workbook.save(content)
        content.seek(0)
        upload = UploadFile(filename='msk47.xlsx', file=content)
        try:
            imported = await self.api.api_import(upload)
        finally:
            await upload.close()
        points = [Point(x=row['X (м)'], y=row['Y (м)']) for row in imported['rows_preview']]
        result = self.convert(code, 'EPSG:4326', points)
        self.assertEqual(len(result.points), 100)
        self.assertAlmostEqual(result.points[-1].lat, original[-1].lat, places=6)
        for kind in ('catalog', 'statement'):
            response = self.api.api_export(kind, self.api.ExportRequest(
                crs_from=code, crs_to='EPSG:4326', rows_in=points, rows_out=result.points,
                format_out='deg_min_sec'))
            data = b''.join([chunk async for chunk in response.body_iterator])
            book = load_workbook(io.BytesIO(data), read_only=True)
            try:
                self.assertEqual(book.active.max_row, 102)
                self.assertEqual(book.active.cell(102, 1).value, 100)
                self.assertTrue(any('°' in str(c) for c in next(book.active.iter_rows(min_row=102, values_only=True))))
            finally:
                book.close()
