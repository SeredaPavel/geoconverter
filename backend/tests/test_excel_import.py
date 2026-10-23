"""Regression checks for Excel import, conversion and both XLSX exports.

Run from backend: python -B -m unittest discover -s tests -v
"""

import importlib
import io
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook

from app.schemas import ConvertRequest, Point


class ExcelImportTests(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        cls.uploads = tempfile.TemporaryDirectory()
        with patch.dict(os.environ, {"UPLOAD_DIR": cls.uploads.name}):
            cls.api = importlib.import_module("app.main")

    @classmethod
    def tearDownClass(cls):
        cls.uploads.cleanup()

    def setUp(self):
        self.upload_patch = patch.object(self.api, "UPLOAD_DIR", Path(self.uploads.name))
        self.upload_patch.start()
        self.addCleanup(self.upload_patch.stop)

    async def import_rows(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["First coordinate", "Second coordinate"])
        for row in rows:
            sheet.append(row)
        content = io.BytesIO()
        workbook.save(content)
        content.seek(0)
        upload = UploadFile(filename="coordinates.xlsx", file=content)
        try:
            return await self.api.api_import(upload)
        finally:
            await upload.close()

    async def test_import_boundaries(self):
        for count in (0, 1, 10, 11, 99, 100, 101, 150):
            with self.subTest(count=count):
                rows = [(55 + i / 10000, 37 + i / 10000) for i in range(count)]
                result = await self.import_rows(rows)
                self.assertEqual(result["rows_total"], count)
                self.assertEqual(result["rows_limit"], 100)
                imported = result["rows_preview"]
                self.assertEqual(len(imported), min(count, 100))
                for index, row in enumerate(imported):
                    actual = list(row.values())
                    self.assertAlmostEqual(actual[0], rows[index][0])
                    self.assertAlmostEqual(actual[1], rows[index][1])

    async def check_exports(self, source, target, rows_in, rows_out):
        request = self.api.ExportRequest(
            crs_from=source, crs_to=target, rows_in=rows_in, rows_out=rows_out
        )
        for kind in ("catalog", "statement"):
            with self.subTest(kind=kind, source=source):
                response = self.api.api_export(kind, request)
                content = b"".join([chunk async for chunk in response.body_iterator])
                workbook = load_workbook(io.BytesIO(content), read_only=True)
                try:
                    sheet = workbook.active
                    self.assertEqual(sheet.max_row, 102)  # Title + headers + 100 points.
                    self.assertEqual(sheet.cell(102, 1).value, 100)
                    fields = ("lat", "lon", "x", "y")
                    expected = []
                    if kind == "statement":
                        expected.extend(getattr(rows_in[-1], f) for f in fields
                                        if getattr(rows_in[-1], f) is not None)
                    expected.extend(getattr(rows_out[-1], f) for f in fields
                                    if getattr(rows_out[-1], f) is not None)
                    for column, value in enumerate(expected, start=2):
                        self.assertAlmostEqual(sheet.cell(102, column).value, value)
                finally:
                    workbook.close()

    async def test_100_geographic_points_convert_and_export(self):
        imported = await self.import_rows([(55 + i / 10000, 37 + i / 10000)
                                           for i in range(100)])
        points = [Point(lat=a, lon=b) for a, b in
                  (row.values() for row in imported["rows_preview"])]
        result = self.api.api_convert(ConvertRequest(
            crs_from="EPSG:4326", crs_to="SK42_GEOG", points=points
        ))
        self.assertEqual(len(result.points), 100)
        await self.check_exports("EPSG:4326", "SK42_GEOG", points, result.points)

    async def test_100_projected_points_convert_and_export(self):
        original = [Point(lat=55 + i / 10000, lon=37 + i / 10000)
                    for i in range(100)]
        projected = self.api.api_convert(ConvertRequest(
            crs_from="EPSG:4326", crs_to="GK_SK42", points=original
        )).points
        # Match the UI, templates and API: easting (x), then northing (y).
        imported = await self.import_rows([(p.x, p.y) for p in projected])
        points = [Point(x=a, y=b) for a, b in
                  (row.values() for row in imported["rows_preview"])]
        result = self.api.api_convert(ConvertRequest(
            crs_from="GK_SK42", crs_to="EPSG:4326", points=points
        ))
        self.assertEqual(len(result.points), 100)
        for expected, actual in zip(original, result.points):
            self.assertAlmostEqual(actual.lat, expected.lat, places=6)
            self.assertAlmostEqual(actual.lon, expected.lon, places=6)
        await self.check_exports("GK_SK42", "EPSG:4326", points, result.points)


if __name__ == "__main__":
    unittest.main()
