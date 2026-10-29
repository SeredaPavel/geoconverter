import importlib
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook

from app.schemas import ConvertRequest, Point
from app.sk63_registry import READY_CRS, SK63_ITEMS, VERIFIED_DEFINITIONS, CATALOG


class Sk63Tests(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        cls.uploads = tempfile.TemporaryDirectory()
        with patch.dict(os.environ, {'UPLOAD_DIR': cls.uploads.name}):
            cls.api = importlib.import_module('app.main')

    @classmethod
    def tearDownClass(cls):
        cls.uploads.cleanup()

    def setUp(self):
        upload_patch = patch.object(self.api, 'UPLOAD_DIR', Path(self.uploads.name))
        upload_patch.start()
        self.addCleanup(upload_patch.stop)

    def convert(self, source, target, points):
        return self.api.api_convert(ConvertRequest(crs_from=source, crs_to=target, points=points))

    async def test_registry_ready_and_dev_statuses(self):
        msk = [item for item in self.api.list_crs() if item.category == 'msk']
        self.assertEqual(len(msk), 260)
        self.assertEqual(sum(item.status == 'ready' for item in msk), 256)
        self.assertEqual(sum(item.status == 'dev' for item in msk), 4)
        choices = [item for item in self.api.list_crs() if item.category == 'sk63']
        self.assertEqual(len({item.code for item in choices}), len(choices))
        self.assertEqual(sum(item.status == 'ready' for item in choices), 120)
        self.assertEqual(sum(item.status == 'dev' for item in choices), 4)
        for item in choices:
            self.assertTrue(item.region)
            self.assertTrue(item.description)
            self.assertEqual(item.code in READY_CRS, item.status == 'ready')

    async def test_all_definitions_origin_coordinates_and_wgs_roundtrip(self):
        for code, crs in READY_CRS.items():
            with self.subTest(code=code):
                params = {p.code: p.value for p in crs.coordinate_operation.params}
                # Known projection invariant also catches an X/Y swap.
                projected = self.convert('SK42_GEOG', code,
                    [Point(lat=params['8801'], lon=(params['8802'] + 180) % 360 - 180)]).points[0]
                self.assertAlmostEqual(projected.x, params['8806'], places=6)
                self.assertAlmostEqual(projected.y, params['8807'], places=6)
                area = crs.area_of_use
                original = (Point(lat=(area.south + area.north) / 2, lon=(area.west + area.east) / 2)
                            if area else Point(**{key: VERIFIED_DEFINITIONS[code]['control_points'][0][key]
                                                  for key in ('lat', 'lon')}))
                result = self.convert('EPSG:4326', code, [original])
                restored = self.convert(code, 'EPSG:4326', result.points).points[0]
                self.assertAlmostEqual(restored.lat, original.lat, places=5)
                self.assertAlmostEqual(restored.lon, original.lon, places=5)
                self.assertFalse(result.meta['allow_ballpark'])
                self.assertTrue(result.meta['operations'])

    async def test_independent_catalogue_reference_points(self):
        expected = {f"SK63:GEOPROJ:{system['source_id']}:{zone['source_value']}"
                    for system in CATALOG['sk63'] for zone in system['zones']}
        self.assertEqual(len(expected), 103)
        self.assertEqual(set(VERIFIED_DEFINITIONS), expected)
        for code, definition in VERIFIED_DEFINITIONS.items():
            for point in definition['control_points'] + definition.get('additional_control_points', []):
                with self.subTest(code=code, lat=point['lat']):
                    if 'epsg' in definition:
                        area = READY_CRS[code].area_of_use
                        self.assertLessEqual(area.west, point['lon'])
                        self.assertLessEqual(point['lon'], area.east)
                        self.assertLessEqual(area.south, point['lat'])
                        self.assertLessEqual(point['lat'], area.north)
                    result = self.convert('EPSG:4326', code, [Point(lat=point['lat'], lon=point['lon'])])
                    self.assertAlmostEqual(result.points[0].x, point['easting'], delta=.01)
                    self.assertAlmostEqual(result.points[0].y, point['northing'], delta=.01)
                    self.assertTrue(result.meta['regional_definitions'])
                    restored = self.convert(code, 'EPSG:4326', result.points).points[0]
                    self.assertAlmostEqual(restored.lat, point['lat'], places=6)
                    self.assertAlmostEqual(restored.lon, point['lon'], places=6)

    async def test_v_zone_is_not_a_samara_or_tatarstan_local_system(self):
        source = self.convert('EPSG:4326', 'SK63:GEOPROJ:V:1', [Point(lat=53.2, lon=50.1)])
        samara = self.convert('SK63:GEOPROJ:V:1', 'MSK:GEOPROJ:63:1', source.points)
        tatarstan = self.convert('SK63:GEOPROJ:V:1', 'MSK:GEOPROJ:16:1', source.points)
        self.assertAlmostEqual(source.points[0].x, samara.points[0].x, delta=.001)
        self.assertAlmostEqual(source.points[0].y - samara.points[0].y, 5500000, delta=.001)
        self.assertAlmostEqual(source.points[0].y - tatarstan.points[0].y, 5700000, delta=.001)
        for target in ('SK63:GEOPROJ:V:2', 'MSK:GEOPROJ:63:2', 'MSK:GEOPROJ:mos:1'):
            with self.subTest(target=target):
                other = self.convert('SK63:GEOPROJ:V:1', target, source.points)
                restored = self.convert(target, 'SK63:GEOPROJ:V:1', other.points).points[0]
                self.assertAlmostEqual(restored.x, source.points[0].x, delta=.02)
                self.assertAlmostEqual(restored.y, source.points[0].y, delta=.02)

    async def test_c_and_t_use_verified_numbering_instead_of_rejected_paper(self):
        # geoproj C1 is EPSG C1 (24.95 degrees); paper C1 is 21.95 degrees.
        crs = READY_CRS['SK63:GEOPROJ:C:1']
        parameters = {p.code: p.value for p in crs.coordinate_operation.params}
        self.assertAlmostEqual(parameters['8802'], 24.95)
        for district, zones, first in [('C', range(3, 7), 24.95), ('T', range(1, 5), 37.98333333333)]:
            for zone in zones:
                crs = READY_CRS[f'SK63:GEOPROJ:{district}:{zone}']
                parameters = {p.code: p.value for p in crs.coordinate_operation.params}
                self.assertAlmostEqual(parameters['8802'], first + (zone - 1) * 3)
        for zone in range(1, 5):
            d = VERIFIED_DEFINITIONS[f'SK63:GEOPROJ:I:{zone}']
            self.assertEqual(d['false_northing'], -12900.60)
            self.assertEqual(d['definition_source'], 'https://geobridge.ru/proj/api')
        for zone in range(1, 7):
            d = VERIFIED_DEFINITIONS[f'SK63:GEOPROJ:L:{zone}']
            self.assertAlmostEqual(d['longitude_natural_origin'], 81 + 31 / 60 + (zone - 1) * 6, places=10)

    async def test_dev_unknown_missing_fields_and_invalid_coordinates(self):
        for item in SK63_ITEMS:
            if item.status != 'dev':
                continue
            for source, target in ((item.code, 'EPSG:4326'), ('EPSG:4326', item.code)):
                with self.subTest(code=item.code), self.assertRaises(HTTPException) as error:
                    self.convert(source, target, [])
                self.assertEqual(error.exception.status_code, 400)
        for points in ([Point(lat=55, lon=30)], [Point(x=1e100, y=1e100)]):
            with self.assertRaises(HTTPException):
                self.convert('SK63:EPSG:3351', 'EPSG:4326', points)
        with self.assertRaises(HTTPException):
            self.convert('SK63:UNKNOWN', 'EPSG:4326', [])

    async def test_sk63_to_sk63_and_gk(self):
        start = self.convert('EPSG:4326', 'SK63:EPSG:7828', [Point(lat=48, lon=32)])
        for target in ('SK63:EPSG:7827', 'GK_SK42', 'GK_SK95', 'GK_GSK2011'):
            with self.subTest(target=target):
                result = self.convert('SK63:EPSG:7828', target, start.points)
                back = self.convert(target, 'SK63:EPSG:7828', result.points)
                self.assertAlmostEqual(back.points[0].x, start.points[0].x, places=2)
                self.assertAlmostEqual(back.points[0].y, start.points[0].y, places=2)

    async def test_100_sk63_points_import_convert_and_both_exports(self):
        await self.check_import_convert_export('SK63:GEOPROJ:P:2', 48, 32)

    async def test_new_batch_100_point_import_export_three_and_six_degree_variants(self):
        for code in ['SK63:GEOPROJ:C:3', 'SK63:GEOPROJ:T:1', 'SK63:GEOPROJ:I:2',
                     'SK63:GEOPROJ:L:2', 'SK63:GEOPROJ:W3:2', 'SK63:GEOPROJ:W6:2',
                     'SK63:GEOPROJ:S:13', 'SK63:GEOPROJ:J:5']:
            control = VERIFIED_DEFINITIONS[code]['control_points'][0]
            with self.subTest(code=code):
                await self.check_import_convert_export(code, control['lat'], control['lon'])

    async def check_import_convert_export(self, code, lat, lon):
        original = [Point(lat=lat + i / 10000, lon=lon + i / 10000) for i in range(100)]
        projected = self.convert('EPSG:4326', code, original).points
        workbook = Workbook()
        workbook.active.append(['X (м)', 'Y (м)'])
        for point in projected:
            workbook.active.append([point.x, point.y])
        content = io.BytesIO()
        workbook.save(content)
        content.seek(0)
        upload = UploadFile(filename='sk63.xlsx', file=content)
        try:
            imported = await self.api.api_import(upload)
        finally:
            await upload.close()
        points = [Point(x=row['X (м)'], y=row['Y (м)']) for row in imported['rows_preview']]
        result = self.convert(code, 'EPSG:4326', points)
        self.assertEqual(len(result.points), 100)
        self.assertAlmostEqual(result.points[-1].lat, original[-1].lat, places=5)
        for kind in ('catalog', 'statement'):
            response = self.api.api_export(kind, self.api.ExportRequest(crs_from=code,
                crs_to='EPSG:4326', rows_in=points, rows_out=result.points, format_out='deg_min_sec'))
            data = b''.join([chunk async for chunk in response.body_iterator])
            book = load_workbook(io.BytesIO(data), read_only=True)
            try:
                self.assertEqual(book.active.max_row, 102)
                self.assertEqual(book.active.cell(102, 1).value, 100)
                self.assertTrue(any('°' in str(c) for c in next(book.active.iter_rows(min_row=102, values_only=True))))
            finally:
                book.close()
