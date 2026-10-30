import importlib
import io
import os
import tempfile
import unittest
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from app.conversion_engine import convert_numeric
from app.schemas import ConvertRequest, Point, UTMOptions
from app.utm_registry import get_crs, CRS_CODE


def convert(source, target, points, **options):
    return convert_numeric(ConvertRequest(crs_from=source, crs_to=target, points=points, options=options))


class UtmTests(unittest.TestCase):
    def test_all_120_epsg_origins_and_inverse(self):
        for hemisphere in ['N', 'S']:
            for zone in range(1, 61):
                with self.subTest(zone=zone, hemisphere=hemisphere):
                    settings = {'zone': zone, 'hemisphere': hemisphere}
                    longitude = zone * 6 - 183
                    result = convert('EPSG:4326', CRS_CODE, [Point(lat=0, lon=longitude)], target_utm=settings)
                    self.assertEqual(get_crs(zone, hemisphere).to_epsg(), (32600 if hemisphere == 'N' else 32700) + zone)
                    self.assertAlmostEqual(result.points[0].x, 500000, delta=1e-6)
                    self.assertAlmostEqual(result.points[0].y, 0 if hemisphere == 'N' else 10000000, delta=1e-6)
                    point = Point(lat=55 if hemisphere == 'N' else -44, lon=longitude + .4)
                    result = convert('EPSG:4326', CRS_CODE, [point], target_utm=settings)
                    restored = convert(CRS_CODE, 'EPSG:4326', result.points, source_utm=settings).points[0]
                    self.assertAlmostEqual(restored.lat, point.lat, places=8)
                    self.assertAlmostEqual(restored.lon, point.lon, places=8)

    def test_published_proj_reference_points_both_hemispheres(self):
        # Independent published examples: PROJ 9.3.0 docs/source/operations/projections/utm.rst.
        for lat, lon, zone, hemisphere, east, north in [(56, 12, 32, 'N', 687071.44, 6210141.33),
                                                       (-44, 174, 59, 'S', 740526.32, 5123750.87)]:
            with self.subTest(hemisphere=hemisphere):
                result = convert('EPSG:4326', CRS_CODE, [Point(lat=lat, lon=lon)],
                                 target_utm={'zone': zone, 'hemisphere': hemisphere}).points[0]
                self.assertAlmostEqual(result.x, east, delta=.01)
                self.assertAlmostEqual(result.y, north, delta=.01)

    def test_independent_zones_with_utm_gk_and_msk(self):
        original = convert('EPSG:4326', CRS_CODE, [Point(lat=55.75, lon=37.6)], target_utm={'zone': 37, 'hemisphere': 'N'})
        other = convert(CRS_CODE, CRS_CODE, original.points,
                        source_utm={'zone': 37, 'hemisphere': 'N'}, target_utm={'zone': 36, 'hemisphere': 'N'})
        self.assertNotAlmostEqual(original.points[0].x, other.points[0].x)
        for target, options in [('GK_GSK2011', {'target_gk': {'zone_mode': 3}}),
                                ('MSK:GEOPROJ:50:2', {}), ('SK63:GEOPROJ:X:4', {})]:
            result = convert(CRS_CODE, target, original.points, source_utm={'zone': 37, 'hemisphere': 'N'}, **options)
            reverse = {'source_gk': {'zone_mode': 3}} if target == 'GK_GSK2011' else {}
            restored = convert(target, CRS_CODE, result.points, target_utm={'zone': 37, 'hemisphere': 'N'}, **reverse)
            self.assertAlmostEqual(restored.points[0].x, original.points[0].x, delta=.02)
            self.assertAlmostEqual(restored.points[0].y, original.points[0].y, delta=.02)

    def test_missing_invalid_zone_hemisphere_and_latitude(self):
        for source, target in [(CRS_CODE, 'EPSG:4326'), ('EPSG:4326', CRS_CODE)]:
            with self.assertRaises(HTTPException) as error:
                convert(source, target, [])
            self.assertEqual(error.exception.status_code, 400)
        for settings in [{'zone': 0, 'hemisphere': 'N'}, {'zone': 61, 'hemisphere': 'S'},
                         {'zone': True, 'hemisphere': 'N'}, {'zone': 1.5, 'hemisphere': 'N'},
                         {'zone': 37, 'hemisphere': 'X'}]:
            with self.assertRaises(ValidationError):
                UTMOptions(**settings)
        for latitude, hemisphere in [(85, 'N'), (-81, 'S'), (-1, 'N'), (1, 'S')]:
            with self.subTest(latitude=latitude), self.assertRaises(HTTPException):
                convert('EPSG:4326', CRS_CODE, [Point(lat=latitude, lon=39)],
                        target_utm={'zone': 37, 'hemisphere': hemisphere})
        for latitude, hemisphere in [(84, 'N'), (-80, 'S')]:
            convert('EPSG:4326', CRS_CODE, [Point(lat=latitude, lon=39)],
                    target_utm={'zone': 37, 'hemisphere': hemisphere})
        with self.assertRaises(HTTPException):
            convert(CRS_CODE, 'EPSG:4326', [Point(x=500000, y=-1000)], source_utm={'zone': 37, 'hemisphere': 'N'})


class UtmExportTests(unittest.IsolatedAsyncioTestCase):
    async def test_100_points_both_exports_preserve_zone_hemisphere_and_axes(self):
        with tempfile.TemporaryDirectory() as uploads, patch.dict(os.environ, {'UPLOAD_DIR': uploads}):
            api = importlib.import_module('app.main')
        for zone, hemisphere, latitude in [(37, 'N', 55), (59, 'S', -44)]:
            settings = {'zone': zone, 'hemisphere': hemisphere}
            points = [Point(lat=latitude + i / 10000, lon=(zone * 6 - 183) + .1 + i / 10000) for i in range(100)]
            result = api.convert(ConvertRequest(crs_from='EPSG:4326', crs_to=CRS_CODE, points=points,
                                                 options={'target_utm': settings}))
            self.assertEqual(len(result.points), 100)
            for kind in ['catalog', 'statement']:
                request = api.ExportRequest(crs_from='EPSG:4326', crs_to=CRS_CODE, rows_in=points,
                                            rows_out=result.points, options={'target_utm': settings})
                response = api.api_export(kind, request)
                content = b''.join([part async for part in response.body_iterator])
                workbook = load_workbook(io.BytesIO(content), read_only=True)
                try:
                    self.assertIn(f'зона {zone}', workbook.active.cell(1, 1).value)
                    self.assertIn('северное' if hemisphere == 'N' else 'южное', workbook.active.cell(1, 1).value)
                    self.assertEqual(workbook.active.max_row, 102)
                    values = list(next(workbook.active.iter_rows(min_row=102, values_only=True)))
                    self.assertAlmostEqual(values[-2], result.points[-1].x, delta=1e-6)
                    self.assertAlmostEqual(values[-1], result.points[-1].y, delta=1e-6)
                finally:
                    workbook.close()
