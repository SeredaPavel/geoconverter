import io
import importlib
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from app.angles import format_angle, parse_angle
from app.export_utils import build_catalog_xlsx, build_statement_xlsx
from app.schemas import ConvertOptions, ConvertRequest, Point


class AngleTests(unittest.TestCase):
    def test_input_formats_and_negative_zero_degrees(self):
        for text in ("55.75833333333333", "55,75833333333333°", "55°45.5′", "55 45 30", "55°45′30″"):
            self.assertAlmostEqual(parse_angle(text), 55 + 45 / 60 + 30 / 3600)
        self.assertAlmostEqual(parse_angle("-0°30′00″"), -0.5)
        self.assertAlmostEqual(Point(lat="-0°30′", lon="37°30′").lat, -0.5)

    def test_invalid_angles_and_ranges(self):
        for text in ("", "abc", "55°60′", "55°30′60″", "55.5°30′", "55°30.5′10″", "NaN", "Infinity", "1 2 3 4"):
            with self.subTest(text=text), self.assertRaises(ValueError):
                parse_angle(text)
        for point in ({"lat": 91}, {"lon": 181}, {"lat": float("nan")}, {"x": float("inf")}):
            with self.subTest(point=point), self.assertRaises(ValidationError):
                Point(**point)

    def test_output_formats_rounding_and_carry(self):
        self.assertEqual(format_angle(55.75833333333333, "deg_min"), "55°45.5000′")
        self.assertEqual(format_angle(55.75833333333333, "deg_min_sec"), "55°45′30.00″")
        self.assertEqual(format_angle(-0.5, "deg_min_sec"), "-0°30′00.00″")
        self.assertEqual(format_angle(12 + 59 / 60 + 59.99999 / 3600, "deg_min_sec"), "13°00′00.00″")

    def test_excel_formats_and_axis_order(self):
        rows = [Point(lat="55°45′30″", lon=37.5), Point(x=7500000, y=6100000)]
        for content, expected in (
            (build_catalog_xlsx(rows[:1], "WGS", "WGS", "deg_min_sec"), [1, "55°45′30.00″", "37°30′00.00″"]),
            (build_statement_xlsx(rows[:1], rows[:1], "WGS", "WGS", "deg_min", "deg_min_sec"),
             [1, "55°45.5000′", "37°30.0000′", "55°45′30.00″", "37°30′00.00″"]),
            (build_catalog_xlsx(rows[1:], "WGS", "GK", "deg_min_sec"), [1, 7500000, 6100000]),
        ):
            workbook = load_workbook(io.BytesIO(content), read_only=True)
            try:
                self.assertEqual(list(next(workbook.active.iter_rows(min_row=3, values_only=True))), expected)
            finally:
                workbook.close()


class CoordinateOptionTests(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        cls.uploads = tempfile.TemporaryDirectory()
        with patch.dict(os.environ, {"UPLOAD_DIR": cls.uploads.name}):
            cls.api = importlib.import_module("app.main")

    @classmethod
    def tearDownClass(cls):
        cls.uploads.cleanup()

    def setUp(self):
        upload_patch = patch.object(self.api, "UPLOAD_DIR", Path(self.uploads.name))
        upload_patch.start()
        self.addCleanup(upload_patch.stop)

    async def test_labelled_excel_axes_are_reordered(self):
        from openpyxl import Workbook
        from fastapi import UploadFile
        for headers, values in ((["Y (м)", "X (м)"], [6100000, 7500000]),
                                (["X (м)", "Y (м)"], [7500000, 6100000]),
                                (["У", "Х"], [6100000, 7500000]),
                                (["Долгота", "Широта"], [37.5, 55.5])):
            with self.subTest(headers=headers):
                workbook = Workbook()
                workbook.active.append(headers)
                workbook.active.append(values)
                content = io.BytesIO()
                workbook.save(content)
                content.seek(0)
                upload = UploadFile(file=content, filename="test.xlsx")
                try:
                    result = await self.api.api_import(upload)
                finally:
                    await upload.close()
                expected = [55.5, 37.5] if headers[0] == "Долгота" else [7500000, 6100000]
                self.assertEqual([result["rows_preview"][0][c] for c in result["columns"]], expected)

    async def test_api_angular_text_and_formatted_output(self):
        result = self.api.convert(ConvertRequest(
            crs_from="EPSG:4326", crs_to="EPSG:4326", format_in="deg_min_sec", format_out="deg_min",
            points=[Point(lat="55°45′30″", lon="37°30′00″")],
        ))
        self.assertAlmostEqual(result.points[0].lat, 55.7583333333)
        self.assertEqual(result.meta["formatted_points"][0], {"lat": "55°45.5000′", "lon": "37°30.0000′"})
        request = self.api.ExportRequest(crs_from="EPSG:4326", crs_to="EPSG:4326",
            rows_in=result.points, rows_out=result.points, format_in="deg_min", format_out="deg_min_sec")
        for kind in ("catalog", "statement"):
            response = self.api.api_export(kind, request)
            content = b"".join([chunk async for chunk in response.body_iterator])
            workbook = load_workbook(io.BytesIO(content), read_only=True)
            try:
                expected = "55°45′30.00″" if kind == "catalog" else "55°45.5000′"
                self.assertEqual(workbook.active.cell(3, 2).value, expected)
            finally:
                workbook.close()

    async def test_manual_zone_and_both_direction_validation(self):
        for target in ("GK_SK42", "GK_SK95", "GK_GSK2011"):
            with self.subTest(target=target):
                automatic = self.api.convert(ConvertRequest(crs_from="EPSG:4326", crs_to=target,
                    points=[Point(lat=55, lon=37)]))
                manual = self.api.convert(ConvertRequest(crs_from="EPSG:4326", crs_to=target,
                    points=[Point(lat=55, lon=37)], options=ConvertOptions(auto_zone=False, zone=8)))
                self.assertEqual(automatic.meta["used_zones"], [7])
                self.assertEqual(manual.meta["used_zones"], [8])
                self.assertNotAlmostEqual(manual.points[0].x, automatic.points[0].x)
                restored = self.api.convert(ConvertRequest(crs_from=target, crs_to="EPSG:4326",
                    points=manual.points, options=ConvertOptions(auto_zone=False, zone=8)))
                self.assertAlmostEqual(restored.points[0].lat, 55, places=6)
                for source, destination, points, options in (
                    ("EPSG:4326", target, [Point(lat=55, lon=37)], ConvertOptions(auto_zone=False)),
                    (target, "EPSG:4326", manual.points, ConvertOptions(auto_zone=False)),
                    ("EPSG:4326", target, [Point(lat=55, lon=37)], ConvertOptions(auto_zone=False, zone=33)),
                    (target, "EPSG:4326", manual.points, ConvertOptions(auto_zone=False, zone=3)),
                ):
                    with self.assertRaises(HTTPException) as error:
                        self.api.convert(ConvertRequest(crs_from=source, crs_to=destination,
                            points=points, options=options))
                    self.assertEqual(error.exception.status_code, 400)

    async def test_existing_templates_match_import_order(self):
        from pathlib import Path
        from fastapi import UploadFile
        path = Path(__file__).resolve().parents[2] / "frontend/public/templates/template_xy.xlsx"
        workbook = load_workbook(path, read_only=True)
        try:
            headers = next(workbook.active.iter_rows(values_only=True))
            self.assertEqual(headers[:2], ("X (м)", "Y (м)"))
        finally:
            workbook.close()
        upload = UploadFile(filename=path.name, file=io.BytesIO(path.read_bytes()))
        try:
            result = await self.api.api_import(upload)
            self.assertEqual(result["columns"], ["X (м)", "Y (м)"])
        finally:
            await upload.close()
